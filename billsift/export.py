"""Excel and CSV export, with a GST summary a CA can use directly."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BILL_COLS = [
    "id", "invoice_date", "vendor_name", "vendor_gstin", "invoice_number", "category",
    "taxable_amount", "gst_rate", "cgst", "sgst", "igst", "total_gst", "total_amount",
    "status", "flags", "source_file",
]
MONEY_COLS = ["taxable_amount", "cgst", "sgst", "igst", "total_gst", "total_amount"]


def financial_year(iso_date: str) -> str:
    try:
        d = datetime.strptime(iso_date, "%Y-%m-%d")
    except (ValueError, TypeError):
        return "Unknown"
    start = d.year if d.month >= 4 else d.year - 1
    return f"FY {start}-{str(start + 1)[-2:]}"


def quarter(iso_date: str) -> str:
    try:
        m = datetime.strptime(iso_date, "%Y-%m-%d").month
    except (ValueError, TypeError):
        return "Unknown"
    return "Q" + str({4: 1, 5: 1, 6: 1, 7: 2, 8: 2, 9: 2, 10: 3, 11: 3, 12: 3, 1: 4, 2: 4, 3: 4}[m])


def bills_dataframe(rows: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(rows)
    if df.empty:
        return pd.DataFrame(columns=BILL_COLS)
    df["total_gst"] = (df["cgst"] + df["sgst"] + df["igst"]).round(2)
    df["flags"] = df["flags"].apply(lambda f: "; ".join(f) if isinstance(f, list) else f)
    return df[BILL_COLS]


def summaries(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    if df.empty:
        empty = pd.DataFrame()
        return {"GST Summary": empty, "By Category": empty, "Needs Review": empty}
    work = df.copy()
    work["FY"] = work["invoice_date"].apply(financial_year)
    work["Quarter"] = work["invoice_date"].apply(quarter)
    work["Month"] = work["invoice_date"].str[:7].replace("", "Unknown")

    agg = {"id": "count", **{c: "sum" for c in MONEY_COLS}}
    gst = (work.groupby(["FY", "Quarter", "Month"], as_index=False).agg(agg)
           .rename(columns={"id": "bills"}).round(2))
    cat = (work.assign(category=work["category"].replace("", "Uncategorised"))
           .groupby("category", as_index=False).agg(agg).rename(columns={"id": "bills"}).round(2))
    review = df[df["status"] == "review"][
        ["id", "invoice_date", "vendor_name", "invoice_number", "total_amount", "flags", "source_file"]
    ]
    return {"GST Summary": gst, "By Category": cat, "Needs Review": review}


def _style(ws) -> None:
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    headers = [c.value for c in ws[1]]
    for name in MONEY_COLS:
        if name in headers:
            idx = headers.index(name) + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=idx).number_format = "#,##0.00"


def build_workbook(rows: list[dict]) -> bytes:
    df = bills_dataframe(rows)
    sheets = {"Bills": df, **summaries(df)}
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, frame in sheets.items():
            frame.to_excel(xw, sheet_name=name, index=False)
            _style(xw.sheets[name])
    return buf.getvalue()


def build_csv(rows: list[dict]) -> bytes:
    return bills_dataframe(rows).to_csv(index=False).encode("utf-8-sig")
